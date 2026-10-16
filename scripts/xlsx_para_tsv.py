#!/usr/bin/env python3
"""Converte cada aba de uma planilha .xlsx em um TSV de texto simples.

Uso: python scripts/xlsx_para_tsv.py <planilha.xlsx> --prefixo <nome> --saida-dir <pasta>

Gera <pasta>/<prefixo>-<aba>.tsv por aba (nome da aba normalizado: minúsculas,
sem acento, não alfanumérico vira "-"). Cada linha começa com "L<n>", o número
da linha na planilha; as células vêm separadas por tabulação, as quebras de
linha dentro da célula viram " ⏎ " e células vazias ficam vazias. As fórmulas
são mantidas como texto (a planilha não é recalculada). Linhas totalmente
vazias são puladas. Se a aba tiver comentários de célula, gera também
<prefixo>-<aba>-comentarios.tsv (referência da célula, texto; o aviso padrão
do Excel sobre comentários encadeados é removido). Requer openpyxl.
"""
from __future__ import annotations

import argparse
import re
import sys
import unicodedata
from pathlib import Path


def normalizar(nome: str) -> str:
    s = unicodedata.normalize("NFKD", nome).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")


def celula(v) -> str:
    if v is None:
        return ""
    t = str(v).replace("\r\n", "\n").replace("\r", "\n").replace("\t", " ")
    return t.replace("\n", " ⏎ ").rstrip()


AVISO_EXCEL = re.compile(r"^\[Threaded comment\].*?Comment:\s*", re.S)


def texto_comentario(c) -> str:
    return celula(AVISO_EXCEL.sub("", c.text or "").strip())


def converter(xlsx: Path, prefixo: str, saida: Path) -> list[tuple[str, int, Path]]:
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, data_only=False)
    saida.mkdir(parents=True, exist_ok=True)
    feitos = []
    for ws in wb.worksheets:
        destino = saida / f"{prefixo}-{normalizar(ws.title)}.tsv"
        n = 0
        with destino.open("w", encoding="utf-8") as f:
            for i, linha in enumerate(ws.iter_rows(values_only=True), start=1):
                vals = [celula(v) for v in linha]
                while vals and vals[-1] == "":
                    vals.pop()
                if not vals:
                    continue
                f.write(f"L{i}\t" + "\t".join(vals) + "\n")
                n += 1
        feitos.append((ws.title, n, destino))
        coment = [(c.coordinate, texto_comentario(c.comment))
                  for linha in ws.iter_rows() for c in linha if c.comment]
        if coment:
            dest_c = saida / f"{prefixo}-{normalizar(ws.title)}-comentarios.tsv"
            with dest_c.open("w", encoding="utf-8") as f:
                for ref, txt in coment:
                    f.write(f"{ref}\t{txt}\n")
            feitos.append((ws.title + " (comentários)", len(coment), dest_c))
    return feitos


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("xlsx", type=Path)
    ap.add_argument("--prefixo", required=True)
    ap.add_argument("--saida-dir", type=Path, required=True)
    a = ap.parse_args(argv)
    for aba, n, destino in converter(a.xlsx, a.prefixo, a.saida_dir):
        print(f"{aba!r}: {n} linhas -> {destino}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
