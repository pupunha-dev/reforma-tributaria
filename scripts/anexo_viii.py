#!/usr/bin/env python3
"""Expande a aba "tabela geral" do Anexo VIII (correlação item da NBS x cIndOp x cClassTrib).

Uso: python scripts/anexo_viii.py <anexo-viii.xlsx> --saida <expandida.tsv> [--validar]

A planilha usa células mescladas: um mesmo item da LC 116 cobre várias NBS, e o
cIndOp (com as colunas "P/S onerosa" e "adquirente exterior"), o local de
incidência e os cClassTrib se estendem por um número diferente de linhas. Este
script "desmescla" cada intervalo (repete o valor do canto superior esquerdo em
todas as linhas do intervalo) e grava uma linha por linha da planilha, com as
colunas: linha, item, descricao_item, nbs, descricao_nbs, ps_onerosa,
adq_exterior, indop, local_incidencia_ibs, cclasstrib, nome_cclasstrib.

Com --validar, confere os códigos contra o Anexo VII (cIndOp) e a tabela
oficial de cClassTrib do Portal dos DF-e (texto/cclasstrib-svrs.json), se
existirem na pasta de textos (--textos). Requer openpyxl.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

COLUNAS = ["linha", "item", "descricao_item", "nbs", "descricao_nbs", "ps_onerosa",
           "adq_exterior", "indop", "local_incidencia_ibs", "cclasstrib", "nome_cclasstrib"]


def limpar(v) -> str:
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v).replace("\n", " ")).strip()


def expandir(xlsx: Path, aba: str = "tabela geral") -> list[list[str]]:
    import openpyxl
    wb = openpyxl.load_workbook(xlsx)
    ws = wb[aba]
    valores = {}
    for linha in ws.iter_rows():
        for c in linha:
            if c.value is not None:
                valores[(c.row, c.column)] = limpar(c.value)
    for rng in ws.merged_cells.ranges:
        topo = valores.get((rng.min_row, rng.min_col), "")
        for r in range(rng.min_row, rng.max_row + 1):
            for col in range(rng.min_col, rng.max_col + 1):
                valores[(r, col)] = topo
    saida = []
    for r in range(2, ws.max_row + 1):
        vals = [valores.get((r, c), "") for c in range(1, 11)]
        if not any(vals):
            continue
        saida.append([str(r)] + vals)
    return saida


def validar(linhas: list[list[str]], textos: Path) -> list[str]:
    avisos = []
    idx = {c: i for i, c in enumerate(COLUNAS)}
    padrao = {"item": r"\d\d\.\d\d(\.\d\d)?", "nbs": r"\d\.\d{4}\.\d\d\.\d\d", "indop": r"\d{6}", "cclasstrib": r"\d{6}"}
    for ln in linhas:
        for campo, rx in padrao.items():
            v = ln[idx[campo]]
            if v and not re.fullmatch(rx, v):
                avisos.append(f"linha {ln[0]}: {campo} fora do padrão: {v!r}")
    indops = {ln[idx["indop"]] for ln in linhas if ln[idx["indop"]]}
    f7 = textos / "anexo-vii-indop-cindop-public.tsv"
    if f7.exists():
        conhecidos = set(re.findall(r"(?<!\d)(\d{6})(?!\d)", f7.read_text(encoding="utf-8")))
        for c in sorted(indops - conhecidos):
            avisos.append(f"cIndOp {c} não consta no Anexo VII")
    fj = textos / "cclasstrib-svrs.json"
    if fj.exists():
        dados = json.loads(fj.read_text(encoding="utf-8"))
        nomes = {}
        for cst in dados:
            for cl in cst.get("ClassificacoesTributarias", []):
                nomes[cl["CodClassTrib"]] = cl.get("NomeClassTrib", "")
        usados = {ln[idx["cclasstrib"]]: ln[idx["nome_cclasstrib"]] for ln in linhas if ln[idx["cclasstrib"]]}
        for c in sorted(set(usados) - set(nomes)):
            avisos.append(f"cClassTrib {c} não consta na tabela oficial (SVRS)")
        # Os nomes do Anexo VIII são rótulos abreviados dos nomes oficiais: não é erro.
    return avisos


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("xlsx", type=Path)
    ap.add_argument("--saida", type=Path, required=True)
    ap.add_argument("--validar", action="store_true")
    ap.add_argument("--textos", type=Path, default=Path("fontes/documentos-fiscais/texto"))
    a = ap.parse_args(argv)
    linhas = expandir(a.xlsx)
    a.saida.parent.mkdir(parents=True, exist_ok=True)
    with a.saida.open("w", encoding="utf-8") as f:
        f.write("\t".join(COLUNAS) + "\n")
        for ln in linhas:
            f.write("\t".join(ln) + "\n")
    print(f"{len(linhas)} linhas -> {a.saida}")
    if a.validar:
        av = validar(linhas, a.textos)
        for x in av:
            print("AVISO:", x)
        print(f"{len(av)} aviso(s)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
